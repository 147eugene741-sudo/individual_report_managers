import os
import sys
import tempfile
import unittest

import openpyxl

TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(TESTS_DIR)
if PROJECT_DIR not in sys.path:
    sys.path.insert(0, PROJECT_DIR)

from backend.parser import AmbiguousAssessmentTypeError, ExcelParser
import main as report_main


def write_form(path, evaluated, survey_name, scores, include_instruction=False):
    """
    scores: dict with keys self/manager/colleagues/subordinates and int or None values.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'ID опроса'
    ws['B1'] = '1'
    ws['A2'] = 'Оцениваемый'
    ws['B2'] = evaluated
    ws['A3'] = 'Дата завершения'
    ws['B3'] = '2024-01-01'
    ws['A4'] = 'Имя опроса'
    ws['B4'] = survey_name
    ws['A5'] = 'Компетенции'
    ws['B5'] = 'Самооценка'
    ws['C5'] = 'Руководитель'
    ws['D5'] = 'Коллеги'
    ws['E5'] = 'Подчиненные'

    row = 6
    if include_instruction:
        ws.cell(row=row, column=1, value='софты инструкция')
        row += 1

    ws.cell(row=row, column=1, value='Ориентация на результат')
    ws.cell(row=row, column=2, value=scores.get('self'))
    ws.cell(row=row, column=3, value=scores.get('manager'))
    ws.cell(row=row, column=4, value=scores.get('colleagues'))
    ws.cell(row=row, column=5, value=scores.get('subordinates'))
    wb.save(path)
    wb.close()


def detect_type(path):
    with ExcelParser(path) as parser:
        header_row = parser._find_row_with_text(1, 'Компетенции')
        return parser._detect_assessment_type(header_row)


class DetectAssessmentTypeTest(unittest.TestCase):
    def setUp(self):
        self._tmpdir = tempfile.TemporaryDirectory()
        self.dir = self._tmpdir.name

    def tearDown(self):
        self._tmpdir.cleanup()

    def _path(self, name):
        return os.path.join(self.dir, name)

    def test_self_only(self):
        path = self._path('self.xlsx')
        write_form(path, 'Иванов Иван', 'Самооценка', {'self': 3})
        self.assertEqual(detect_type(path), 'self')

    def test_manager_only(self):
        path = self._path('manager.xlsx')
        write_form(path, 'Иванов Иван', 'Оценка руководителя', {'manager': 4})
        self.assertEqual(detect_type(path), 'manager')

    def test_mixed_self_and_colleagues_is_colleagues(self):
        path = self._path('mixed_colleagues.xlsx')
        write_form(path, 'Иванов Иван', 'Самооценка', {'self': 2, 'colleagues': 4})
        self.assertEqual(detect_type(path), 'colleagues')

    def test_mixed_self_and_manager_is_manager(self):
        path = self._path('mixed_manager.xlsx')
        write_form(path, 'Иванов Иван', 'Самооценка', {'self': 3, 'manager': 1})
        self.assertEqual(detect_type(path), 'manager')

    def test_mixed_self_and_subordinates_is_subordinates(self):
        path = self._path('mixed_sub.xlsx')
        write_form(path, 'Иванов Иван', 'Самооценка', {'self': 3, 'subordinates': 2})
        self.assertEqual(detect_type(path), 'subordinates')

    def test_multiple_non_self_roles_are_ambiguous(self):
        path = self._path('ambiguous.xlsx')
        write_form(
            path,
            'Иванов Иван',
            'Оценка',
            {'self': 3, 'manager': 2, 'colleagues': 4},
        )
        with self.assertRaises(AmbiguousAssessmentTypeError) as ctx:
            detect_type(path)
        self.assertIn('self', ctx.exception.found_types)
        self.assertIn('manager', ctx.exception.found_types)
        self.assertIn('colleagues', ctx.exception.found_types)

    def test_manager_and_colleagues_without_self_are_ambiguous(self):
        path = self._path('two_roles.xlsx')
        write_form(path, 'Иванов Иван', 'Оценка', {'manager': 2, 'colleagues': 3})
        with self.assertRaises(AmbiguousAssessmentTypeError):
            detect_type(path)

    def test_fallback_by_survey_name(self):
        path = self._path('fallback.xlsx')
        write_form(path, 'Иванов Иван', 'Оценка коллег', {})
        self.assertEqual(detect_type(path), 'colleagues')

    def test_parse_mixed_form_reads_only_role_column(self):
        path = self._path('mixed_parse.xlsx')
        write_form(
            path,
            'Иванов Иван',
            'Самооценка',
            {'self': 2, 'colleagues': 4},
            include_instruction=True,
        )
        with ExcelParser(path) as parser:
            data = parser.parse()
        self.assertEqual(parser.assessment_type, 'colleagues')
        self.assertEqual(data['competencies']['soft']['scores'], [4])


class GroupFilesByEmployeeTest(unittest.TestCase):
    def setUp(self):
        self._tmpdir = tempfile.TemporaryDirectory()
        self.dir = self._tmpdir.name

    def tearDown(self):
        self._tmpdir.cleanup()

    def _path(self, name):
        return os.path.join(self.dir, name)

    def test_clean_self_and_dirty_colleagues(self):
        write_form(self._path('self.xlsx'), 'Петров Пётр', 'Самооценка', {'self': 3})
        write_form(self._path('manager.xlsx'), 'Петров Пётр', 'Руководитель', {'manager': 2})
        write_form(
            self._path('dirty_colleagues.xlsx'),
            'Петров Пётр',
            'Самооценка',
            {'self': 3, 'colleagues': 4},
        )

        groups, skipped = report_main.group_files_by_employee(self.dir)
        self.assertEqual(skipped, [])
        self.assertIn('Петров Пётр', groups)
        files = groups['Петров Пётр']
        self.assertTrue(files['self'].endswith('self.xlsx'))
        self.assertTrue(files['colleagues'].endswith('dirty_colleagues.xlsx'))
        self.assertTrue(files['manager'].endswith('manager.xlsx'))

    def test_ambiguous_file_is_skipped_and_reported(self):
        write_form(self._path('self.xlsx'), 'Сидоров Sid', 'Самооценка', {'self': 3})
        write_form(self._path('manager.xlsx'), 'Сидоров Sid', 'Руководитель', {'manager': 2})
        write_form(self._path('colleagues.xlsx'), 'Сидоров Sid', 'Коллеги', {'colleagues': 4})
        write_form(
            self._path('ambiguous.xlsx'),
            'Сидоров Sid',
            'Оценка',
            {'manager': 1, 'colleagues': 2},
        )

        groups, skipped = report_main.group_files_by_employee(self.dir)
        self.assertEqual(len(skipped), 1)
        self.assertEqual(skipped[0]['file'], 'ambiguous.xlsx')
        self.assertEqual(skipped[0]['employee'], 'Сидоров Sid')
        self.assertIn('Сидоров Sid', groups)
        self.assertTrue(groups['Сидоров Sid']['colleagues'].endswith('colleagues.xlsx'))

    def test_only_ambiguous_non_self_file_skips_employee(self):
        write_form(self._path('self.xlsx'), 'Козлов К', 'Самооценка', {'self': 3})
        write_form(
            self._path('both.xlsx'),
            'Козлов К',
            'Оценка',
            {'manager': 2, 'colleagues': 4},
        )

        groups, skipped = report_main.group_files_by_employee(self.dir)
        self.assertEqual(groups, {})
        self.assertEqual(len(skipped), 1)
        self.assertEqual(skipped[0]['file'], 'both.xlsx')


if __name__ == '__main__':
    unittest.main()
